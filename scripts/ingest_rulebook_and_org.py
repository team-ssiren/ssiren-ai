"""SQLite 적재 스크립트 (멱등).

- docs/rules/database/<MAJOR>/<MINOR>.md → assignment_guide (파일 전체를 평문 저장)
- data/seed/분당기관부서.xlsx → agency / department

실행: uv run python scripts/ingest_rulebook_and_org.py
소스 파일은 읽기만, .db 만 갱신한다.
"""

from __future__ import annotations

import os
import sys

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from pathlib import Path  # noqa: E402

from openpyxl import load_workbook  # noqa: E402

from app.db.connection import get_connection, run_schema  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
RULEBOOK_DIR = REPO_ROOT / "docs" / "rules" / "database"
SEED_XLSX = REPO_ROOT / "data" / "seed" / "분당기관부서.xlsx"
REGION = "BUNDANG"


def ingest_rulebook(conn) -> int:
    conn.execute("DELETE FROM assignment_guide")
    count = 0
    for md_path in sorted(RULEBOOK_DIR.glob("*/*.md")):
        major_code = md_path.parent.name
        minor_code = md_path.stem
        guide_text = md_path.read_text(encoding="utf-8")
        conn.execute(
            "INSERT OR REPLACE INTO assignment_guide(major_code, minor_code, guide_text) "
            "VALUES (?, ?, ?)",
            (major_code, minor_code, guide_text),
        )
        count += 1
    return count


def _agency_type(full_name: str) -> str | None:
    """전체기관명으로 4값(지자체/경찰/소방/보건) 유형 도출. 세무서 등은 None(스킵)."""
    if "소방" in full_name or "119" in full_name:
        return "소방"
    if "경찰" in full_name or "지구대" in full_name or "파출소" in full_name:
        return "경찰"
    if "보건소" in full_name or "보건지소" in full_name:
        return "보건"
    if "세무서" in full_name:
        return None
    if "구청" in full_name or "동사무소" in full_name:
        return "지자체"
    return "지자체"


def _clean(value) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    return s or None


def ingest_org(conn) -> tuple[int, int]:
    conn.execute("DELETE FROM department")
    conn.execute("DELETE FROM agency")

    wb = load_workbook(SEED_XLSX, read_only=True, data_only=True)
    ws = wb.active
    rows = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows)]
    idx = {name: i for i, name in enumerate(header)}

    def col(row, name):
        return row[idx[name]] if name in idx and idx[name] < len(row) else None

    agency_ids: dict[str, int] = {}
    agency_count = dept_count = 0

    for row in rows:
        full_name = _clean(col(row, "전체기관명"))
        dept_name = _clean(col(row, "최하위기관명"))
        if not full_name or not dept_name:
            continue
        atype = _agency_type(full_name)
        if atype is None:
            continue
        phone = _clean(col(row, "전화번호"))

        agency_id = agency_ids.get(full_name)
        if agency_id is None:
            cur = conn.execute(
                "INSERT INTO agency(region_code, agency_type, name, phone) VALUES (?, ?, ?, ?)",
                (REGION, atype, full_name, phone),
            )
            agency_id = cur.lastrowid
            agency_ids[full_name] = agency_id
            agency_count += 1
        elif phone:
            conn.execute(
                "UPDATE agency SET phone = COALESCE(phone, ?) WHERE id = ?", (phone, agency_id)
            )

        conn.execute(
            "INSERT OR IGNORE INTO department(agency_id, region_code, agency_type, name, phone) "
            "VALUES (?, ?, ?, ?, ?)",
            (agency_id, REGION, atype, dept_name, phone),
        )
        dept_count += 1

    wb.close()
    return agency_count, dept_count


def main() -> None:
    run_schema()
    conn = get_connection()
    guides = ingest_rulebook(conn)
    agencies, depts = ingest_org(conn)
    conn.commit()
    print(f"assignment_guide: {guides} rows")
    print(f"agency: {agencies} rows / department: {depts} rows")


if __name__ == "__main__":
    main()
